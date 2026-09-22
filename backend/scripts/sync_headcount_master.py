import openpyxl
import logging
import sys
import os

# Add backend directory to sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from core.config import settings
from supabase import create_client

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("sync_headcount_master")

EXCEL_PATH = "/Users/khwaish/Downloads/Final Data/RLL_User_Onboarding_Template.xlsx"

def sync_headcount_master():
    logger.info(f"Connecting to Supabase Test DB: {settings.SUPABASE_URL}")
    client = create_client(settings.SUPABASE_URL, settings.SUPABASE_KEY)

    logger.info(f"Loading onboarding template from: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb['Users']

    # 1. Fetch existing Roles Map
    roles_res = client.table("roles").select("role_id, role_name").execute()
    role_id_map = {r["role_name"].upper(): str(r["role_id"]) for r in (roles_res.data or []) if r.get("role_id")}
    logger.info(f"Roles in Database: {role_id_map}")

    # Map Excel roles to DB core roles
    # Excel roles: Admin -> ADMIN, Team Leader -> LEADER, ASM / TSM -> TSM, ASE -> ASE
    def resolve_role_name(excel_role: str) -> str:
        r = str(excel_role or "").strip().lower()
        if "admin" in r:
            return "ADMIN"
        elif "leader" in r or "lead" in r:
            return "LEADER"
        elif "tsm" in r or "asm" in r or "manager" in r:
            return "TSM"
        else:
            return "ASE"

    # Read rows from Excel
    excel_records = []
    for r in range(2, ws.max_row + 1):
        fn = ws.cell(r, 1).value
        ln = ws.cell(r, 2).value
        ph = ws.cell(r, 3).value
        em = ws.cell(r, 4).value
        role = ws.cell(r, 5).value
        mgr = ws.cell(r, 6).value

        if not fn or str(fn).strip().lower() in ("none", "null", ""):
            continue

        fn_clean = str(fn).strip()
        ln_clean = str(ln).strip() if ln and str(ln).strip().lower() not in ("none", "null") else ""
        full_name = f"{fn_clean} {ln_clean}".strip()
        email_clean = str(em).strip().lower() if em and str(em).strip().lower() not in ("none", "null") else f"{fn_clean.lower()}.{ln_clean.lower() or 'user'}@rll.co.in"
        phone_clean = str(ph).strip().replace("+91", "").replace(" ", "") if ph else f"9900{r:06d}"
        role_db = resolve_role_name(role)
        mgr_clean = str(mgr).strip() if mgr and str(mgr).strip().lower() not in ("none", "null") else ""

        excel_records.append({
            "row": r,
            "first_name": fn_clean,
            "last_name": ln_clean,
            "full_name": full_name,
            "email": email_clean,
            "phone": phone_clean,
            "role": role_db,
            "raw_role": role,
            "manager": mgr_clean
        })

    logger.info(f"Parsed {len(excel_records)} valid employee records from Excel.")

    # 2. Synchronize public.users
    name_to_user_id = {}
    email_to_user_id = {}

    for rec in excel_records:
        email = rec["email"]
        fn = rec["first_name"]
        ln = rec["last_name"]
        ph = rec["phone"]

        # Check existing user by email or name
        res = client.table("users").select("user_id, email, first_name, last_name").ilike("email", email).execute()
        user_id = None
        if res.data:
            user_id = str(res.data[0]["user_id"])
            # Update user info
            client.table("users").update({
                "first_name": fn,
                "last_name": ln,
                "phone": ph,
                "is_active": True
            }).eq("user_id", user_id).execute()
        else:
            # Check by first_name + last_name
            res_name = client.table("users").select("user_id").ilike("first_name", fn).ilike("last_name", ln).execute()
            if res_name.data:
                user_id = str(res_name.data[0]["user_id"])
                client.table("users").update({
                    "email": email,
                    "phone": ph,
                    "is_active": True
                }).eq("user_id", user_id).execute()
            else:
                # Create user
                ins_res = client.table("users").insert({
                    "email": email,
                    "first_name": fn,
                    "last_name": ln,
                    "phone": ph,
                    "is_active": True
                }).execute()
                if ins_res.data:
                    user_id = str(ins_res.data[0]["user_id"])

        if user_id:
            rec["user_id"] = user_id
            name_to_user_id[rec["full_name"].lower()] = user_id
            name_to_user_id[fn.lower()] = user_id
            email_to_user_id[email] = user_id

    # 3. Synchronize public.user_roles (Exactly 1 active primary role per user)
    logger.info("Synchronizing user_roles table...")
    synced_user_ids = {rec["user_id"] for rec in excel_records if rec.get("user_id")}

    for rec in excel_records:
        u_id = rec.get("user_id")
        role_name = rec["role"]
        r_id = role_id_map.get(role_name)

        if not u_id or not r_id:
            continue

        # Deactivate all existing roles for this user first
        client.table("user_roles").update({"is_active": False}).eq("user_id", u_id).execute()

        # Upsert active role
        client.table("user_roles").upsert({
            "user_id": u_id,
            "role_id": r_id,
            "is_active": True
        }, on_conflict="user_id,role_id").execute()

    # Deactivate duplicate user_roles for users not in roster or deactivate stale dummy users
    all_users = client.table("users").select("user_id, email, first_name, last_name").execute().data or []
    for u in all_users:
        uid = str(u["user_id"])
        if uid not in synced_user_ids:
            # Deactivate stale dummy users auto-created by past sales uploads (e.g. @rll.com duplicates)
            em = str(u.get("email") or "").lower()
            if ".com" in em or "dummy" in em or "unassigned" in em or "test" in em:
                client.table("users").update({"is_active": False}).eq("user_id", uid).execute()
                client.table("user_roles").update({"is_active": False}).eq("user_id", uid).execute()

    # 5. Rebuild public.ase_tsm_mapping (Strictly TSM -> ASE field mappings)
    logger.info("Rebuilding ase_tsm_mapping table strictly from onboarding template...")
    # Clear existing mapping table
    client.table("ase_tsm_mapping").delete().neq("hierarchy_id", "00000000-0000-0000-0000-000000000000").execute()

    tsm_ase_count = 0
    for rec in excel_records:
        u_id = rec.get("user_id")
        role = rec["role"]
        mgr_name = rec["manager"]

        # In ase_tsm_mapping, we ONLY map ASEs whose manager is a TSM/ASM
        if role == "ASE" and mgr_name:
            mgr_id = name_to_user_id.get(mgr_name.lower())
            if not mgr_id:
                parts = mgr_name.split()
                mgr_id = name_to_user_id.get(parts[0].lower())

            if mgr_id:
                # Verify manager is TSM/ASM
                mgr_rec = next((r for r in excel_records if r.get("user_id") == mgr_id), None)
                if mgr_rec and mgr_rec["role"] in ["TSM", "ASM"]:
                    try:
                        client.table("ase_tsm_mapping").insert({
                            "tsm_user_id": mgr_id,
                            "ase_user_id": u_id,
                            "is_active": True,
                            "ase_name": rec["full_name"]
                        }).execute()
                        tsm_ase_count += 1
                    except Exception as e_ins:
                        logger.warning(f"Error inserting mapping for {rec['full_name']} under {mgr_name}: {e_ins}")

    # 6. Rebind user_sales_fact records from inactive user_ids to active user_ids
    logger.info("Rebinding user_sales_fact records from inactive user_ids to active user_ids...")
    try:
        all_u_res = client.table("users").select("user_id, first_name, last_name, email, is_active").execute()
        all_u_list = all_u_res.data or []
        act_u_list = [u for u in all_u_list if u.get("is_active")]

        def _clean_str(s):
            return "".join(c for c in str(s or "").lower() if c.isalnum())

        def _first_word(s):
            parts = str(s or "").strip().lower().replace(".", " ").replace("_", " ").split()
            return _clean_str(parts[0]) if parts else ""

        active_by_full = {_clean_str(f"{u.get('first_name','')} {u.get('last_name','')}"): str(u["user_id"]) for u in act_u_list}
        active_by_first = {_first_word(u.get('first_name','')): str(u["user_id"]) for u in act_u_list if _first_word(u.get('first_name',''))}
        active_by_prefix = {_first_word(str(u.get('email') or '').split('@')[0]): str(u["user_id"]) for u in act_u_list if _first_word(str(u.get('email') or '').split('@')[0])}

        rebound_count = 0
        for u in all_u_list:
            uid = str(u["user_id"])
            if not u.get("is_active"):
                fn = str(u.get("first_name") or "").strip()
                ln = str(u.get("last_name") or "").strip()
                full_name = f"{fn} {ln}".strip()
                email = str(u.get("email") or "")
                prefix = email.split("@")[0]

                act_id = (
                    active_by_full.get(_clean_str(full_name)) or
                    active_by_prefix.get(_first_word(prefix)) or
                    active_by_first.get(_first_word(fn)) or
                    active_by_first.get(_first_word(full_name))
                )
                if act_id and act_id != uid:
                    try:
                        up_res = client.table("user_sales_fact").update({"user_id": act_id}).eq("user_id", uid).execute()
                        rebound_count += len(up_res.data or [])
                    except Exception as e_up:
                        logger.warning(f"Error rebinding user_sales_fact for {uid} -> {act_id}: {e_up}")

        logger.info(f"Rebound {rebound_count} user_sales_fact rows to active user IDs!")
    except Exception as e_reb:
        logger.warning(f"Error in user_sales_fact rebinding: {e_reb}")

    logger.info(f"SUCCESS: Synchronized {len(excel_records)} users, updated roles, and established {tsm_ase_count} TSM->ASE field mappings!")

if __name__ == "__main__":
    sync_headcount_master()
